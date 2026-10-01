"""Automated tests for the CITY LINKS Recovery System (spec section 57).

Run:  python -m pytest tests -v
"""
from __future__ import annotations

import datetime as dt
import shutil
import sys
import warnings
from pathlib import Path

import pytest

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))
warnings.filterwarnings("ignore")

from citylinks import backup as bk  # noqa: E402
from citylinks import importer as imp  # noqa: E402
from citylinks import pdfgen  # noqa: E402
from citylinks import schema as S  # noqa: E402
from citylinks import services as svc  # noqa: E402
from citylinks.cli import App, export_reports, whatsapp_link  # noqa: E402
from citylinks.seed import build_demo, build_real  # noqa: E402
from citylinks.store import Store, WorkbookOpenError, lock_file  # noqa: E402
from citylinks.util import UserError, parse_month, safe_filename, whatsapp_number  # noqa: E402

M1, M2 = "2026-10", "2026-11"


@pytest.fixture
def wb(tmp_path):
    p = tmp_path / "data" / "CITY_LINKS_Recovery.xlsx"
    p.parent.mkdir()
    build_real(p, current_month=M1)
    return p


def S_(p):
    return Store(p)


def add(p, **kw):
    st = S_(p)
    base = {"Customer Name": "Shahid Mahmood", "Location Code": "A1", "VLAN ID": "405", "Due Date": 10,
            "Fee Override": 1300, "Customer Status": "Active"}
    base.update(kw)
    cid = svc.add_customer(st, base, confirm_duplicates=kw.pop("_confirm", None))
    st.save()
    return cid


def gen(p, month=M1, **kw):
    st = S_(p)
    r = svc.generate_month(st, month, **kw)
    st.save()
    return r


def pay(p, cid, amount, month=M1, **kw):
    st = S_(p)
    pid = svc.record_payment(st, cid, month, amount, kw.pop("date", dt.date(2026, 10, 5)), kw.pop("method", "Cash"),
                             kw.pop("collector", "B01"), **kw)
    st.save()
    return pid


def data(p):
    return svc.load(Store(p, check_lock=False))


# ------------------------------------------------------------------ customers
def test_new_customer_gets_permanent_id_and_event(wb):
    cid = add(wb)
    assert cid == "CL-0001"
    assert add(wb, **{"Customer Name": "شاہد محمود", "VLAN ID": "406", "Mobile Number": "03001234567"}) == "CL-0002"
    d = data(wb)
    assert d.customer("CL-0002")["Customer Name"] == "شاہد محمود"         # Urdu kept exactly, not translated
    assert [e["Event Type"] for e in d.events if e["Customer ID"] == "CL-0001"] == ["New Connection"]


def test_prefix_is_configurable(wb):
    st = S_(wb)
    st.set_setting("CustomerIDPrefix", "CTL")
    st.save()
    assert add(wb) == "CTL-0001"


def test_duplicate_customer_id_is_refused(wb):
    add(wb)
    with pytest.raises(UserError, match="already exists"):
        add(wb, **{"Customer ID": "CL-0001", "VLAN ID": "999", "Customer Name": "Other"})


def test_possible_duplicate_needs_confirmation(wb):
    add(wb, **{"Mobile Number": "0300-1234567"})
    with pytest.raises(UserError, match="Possible duplicate"):
        add(wb, **{"Customer Name": "SHAHID  mahmood", "VLAN ID": "500"})       # same name+location, case/space differ
    with pytest.raises(UserError, match="same Mobile"):
        add(wb, **{"Customer Name": "Someone", "VLAN ID": "501", "Mobile Number": "03001234567"})
    cid = add(wb, **{"Customer Name": "Shahid Mahmood", "VLAN ID": "502", "_confirm": lambda d: True})
    assert cid == "CL-0002"                                                   # saved only after confirmation, not merged
    assert len(data(wb).customers) == 2


def test_location_code_A1_and_B1_derive_address(wb):
    a = add(wb)
    b = add(wb, **{"Customer Name": "Imran", "Location Code": "b1", "VLAN ID": "207", "Fee Override": 1500})
    d = data(wb)
    assert (d.customer(a)["Area"], d.customer(a)["Street"], d.customer(a)["Full Address"]) == ("Arif Town", "Gali 1", "Arif Town, Gali 1")
    assert d.customer(b)["Full Address"] == "Arai Colony, Gali 1"
    assert d.customer(b)["Location Code"] == "B1"


def test_invalid_location_code(wb):
    with pytest.raises(UserError, match="Location Code A7 does not exist"):
        add(wb, **{"Location Code": "A7"})


def test_new_area_and_location(wb):
    st = S_(wb)
    svc.add_area(st, "C", "New Area")
    svc.add_location(st, "C1")
    st.save()
    cid = add(wb, **{"Location Code": "C1", "VLAN ID": "700"})
    assert data(wb).customer(cid)["Full Address"] == "New Area, Gali 1"


def test_required_fields_messages(wb):
    with pytest.raises(UserError, match="VLAN ID is missing"):
        add(wb, **{"VLAN ID": ""})
    with pytest.raises(UserError, match="Monthly Fee is missing"):
        add(wb, **{"Fee Override": None, "VLAN ID": "1"})


def test_vlan_text_keeps_leading_zero(wb):
    cid = add(wb, **{"VLAN ID": "0405"})
    assert data(wb).customer(cid)["VLAN ID"] == "0405"


def test_package_fee_and_override(wb):
    a = add(wb, **{"Package": "Standard", "Fee Override": None})
    b = add(wb, **{"Customer Name": "X", "Package": "Standard", "Fee Override": 1200, "VLAN ID": "9"})
    d = data(wb)
    assert d.customer(a)["Monthly Fee"] == 1500 and d.customer(b)["Monthly Fee"] == 1200


# -------------------------------------------------------------------- billing
def test_month_generation_only_active_and_snapshot(wb):
    a = add(wb)
    b = add(wb, **{"Customer Name": "Gone", "VLAN ID": "410", "Customer Status": "Disconnected"})
    r = gen(wb)
    assert r.billed == 1
    d = data(wb)
    bills = [x for x in d.billing if x["Month"] == M1]
    assert [x["Customer ID"] for x in bills] == [a]
    x = bills[0]
    assert (x["Amount Paid"], x["Balance"], x["Payment Status"], x["Billing ID"]) == (0, 1300, "PENDING", f"{M1}-{a}")
    assert d.settings["CurrentMonth"] == M1
    # later fee change does not change October's bill (history preserved)
    st = S_(wb)
    svc.update_customer(st, a, {"Fee Override": 1500})
    st.save()
    gen(wb, M2)
    d = data(wb)
    assert svc.bill_for(d, a, M1)["Monthly Fee"] == 1300
    assert svc.bill_for(d, a, M2)["Monthly Fee"] == 1500
    assert svc.bill_for(d, b, M2) is None


def test_duplicate_month_is_refused(wb):
    add(wb)
    gen(wb)
    with pytest.raises(UserError, match="October 2026 billing already exists"):
        gen(wb)
    assert len(data(wb).billing) == 1


def test_month_with_problem_customer(wb):
    add(wb)
    st = S_(wb)
    ws = st.wb["CUSTOMERS"]
    ws.cell(6, 1, "CL-0099"); ws.cell(6, 2, "Manual Row"); ws.cell(6, 3, "Z9"); ws.cell(6, 8, "1"); ws.cell(6, 17, "Active")
    ws.cell(6, 11, 1000); ws.cell(6, 13, 10)
    st.t(S.CUSTOMERS).table.ref = "A4:V6"
    st.save()
    with pytest.raises(UserError, match="Location Code Z9 does not exist"):
        gen(wb)
    r = gen(wb, skip_invalid=True)
    assert r.billed == 1 and len(r.skipped) == 1


# ------------------------------------------------------------------- payments
def test_full_partial_multiple_and_pending(wb):
    a = add(wb, **{"Fee Override": 1500})
    b = add(wb, **{"Customer Name": "B", "VLAN ID": "2"})
    c = add(wb, **{"Customer Name": "C", "VLAN ID": "3"})
    gen(wb)
    pay(wb, a, 500)
    assert svc.bill_for(data(wb), a, M1)["Payment Status"] == "PARTIAL"
    pay(wb, a, 500)
    pay(wb, a, 500)
    x = svc.bill_for(data(wb), a, M1)
    assert (x["Amount Paid"], x["Balance"], x["Payment Status"], x["Payments Count"]) == (1500, 0, "PAID", 3)
    pay(wb, b, 1300)
    assert svc.bill_for(data(wb), b, M1)["Payment Status"] == "PAID"
    assert svc.bill_for(data(wb), c, M1)["Payment Status"] == "PENDING"
    assert len([p for p in data(wb).payments if p["Customer ID"] == a]) == 3     # individual transactions kept


def test_payment_errors(wb):
    a = add(wb)
    gen(wb)
    with pytest.raises(UserError, match="cannot be negative"):
        pay(wb, a, -5)
    with pytest.raises(UserError, match="Customer CL-0025 was not found"):
        pay(wb, "CL-0025", 100)
    with pytest.raises(UserError, match="more than the remaining"):
        pay(wb, a, 2000)
    pay(wb, a, 2000, allow_overpay=True)
    x = svc.bill_for(data(wb), a, M1)
    assert x["Balance"] == 0 and x["Excess Paid"] == 700          # no negative balance by default


def test_void_payment_keeps_record(wb):
    a = add(wb)
    gen(wb)
    pid = pay(wb, a, 1300)
    st = S_(wb)
    svc.void_payment(st, pid, "entered twice")
    st.save()
    d = data(wb)
    assert svc.bill_for(d, a, M1)["Payment Status"] == "PENDING"
    assert any(p["Payment ID"] == pid and p["Voided"] == "Yes" for p in d.payments)


# ---------------------------------------------------------- filters / print
def _area_setup(wb):
    ids = [add(wb, **{"Customer Name": f"N{i}", "Location Code": loc, "VLAN ID": str(100 + i), "Due Date": due})
           for i, (loc, due) in enumerate([("A1", 10), ("A1", 5), ("A2", 10), ("B1", 10), ("B2", 5)])]
    gen(wb)
    st = S_(wb)
    svc.assign_collector(st, M1, "B01", ids[:3])
    svc.assign_collector(st, M1, "B02", ids[3:])
    st.save()
    pay(wb, ids[0], 1300)
    pay(wb, ids[3], 300, collector="B02")
    return ids


def test_area_street_collector_due_status_filters(wb):
    ids = _area_setup(wb)
    d = data(wb)
    ids_of = lambda bills: [b["Customer ID"] for b in bills]
    assert ids_of(svc.select_bills(d, M1, area="Arif Town")) == ids[:3]
    assert ids_of(svc.select_bills(d, M1, location="A1")) == ids[:2]
    assert ids_of(svc.select_bills(d, M1, collector="Ali")) == ids[:3]
    assert ids_of(svc.select_bills(d, M1, collector="B02")) == ids[3:]
    assert ids_of(svc.select_bills(d, M1, due=5)) == [ids[1], ids[4]]
    assert ids_of(svc.select_bills(d, M1, area="Arif Town", location="A1", status="PENDING", collector="Ali")) == [ids[1]]
    assert ids_of(svc.select_bills(d, M1, status="PARTIAL")) == [ids[3]]
    assert ids_of(svc.select_bills(d, M1, mode="Pending Recovery")) == ids[1:]          # pending + partial
    assert ids_of(svc.select_bills(d, M1, mode="Custom Filtered Recovery", customer_ids=[ids[2]])) == [ids[2]]


def test_reassignment_keeps_history(wb):
    ids = _area_setup(wb)
    st = S_(wb)
    svc.assign_collector(st, M1, "B03", [ids[0]])
    st.save()
    d = data(wb)
    assert svc.bill_for(d, ids[0], M1)["Collector ID"] == "B03"
    rows = [a for a in d.assignments if a["Customer ID"] == ids[0]]
    assert sorted(a["Assignment Status"] for a in rows) == ["Assigned", "Reassigned"]


def test_recovery_list_pdf_english_and_urdu(wb, tmp_path):
    _area_setup(wb)
    d = data(wb)
    bills = svc.select_bills(d, M1, area="Arif Town", status="UNPAID")
    f = {"month": M1, "area": "Arif Town", "location": "All", "collector": "All", "status": "UNPAID"}
    for lang in ("English", "Urdu"):
        p = pdfgen.recovery_list_pdf(bills, f, d.settings, tmp_path, lang)
        assert p.exists() and p.read_bytes()[:4] == b"%PDF"


def test_dashboard_and_cash_reconciliation(wb):
    ids = _area_setup(wb)
    d = data(wb)
    s = svc.dashboard(d, M1, today=dt.date(2026, 10, 5))
    assert (s["customers"], s["billing"], s["recovered"], s["pending"]) == (5, 6500, 1600, 4900)
    assert (s["paid"], s["partial"], s["pending_count"]) == (1, 1, 3)
    assert s["today_collection"] == 1600
    st = S_(wb)
    svc.set_cash_submitted(st, M1, "B01", 1300)
    svc.set_cash_submitted(st, M1, "B02", 200, remarks="Rs 100 to be given tomorrow")
    st.save()
    rec = {r["Collector ID"]: r for r in svc.cash_reconciliation(data(wb), M1)}
    assert (rec["B01"]["Expected Cash"], rec["B01"]["Difference"], rec["B01"]["Status"]) == (1300, 0, "MATCHED")
    assert (rec["B02"]["Expected Cash"], rec["B02"]["Difference"], rec["B02"]["Status"]) == (300, -100, "CASH DIFFERENCE")
    assert rec["B03"]["Status"] == "NOT SUBMITTED"


def test_overdue_definition(wb):
    a = add(wb, **{"Due Date": 31})
    gen(wb, "2026-09")
    b = svc.bill_for(data(wb), a, "2026-09")
    assert not svc.is_overdue(b, today=dt.date(2026, 9, 30))       # due 31 -> last day (30 Sep)
    assert svc.is_overdue(b, today=dt.date(2026, 10, 1))


# ------------------------------------------------------- history / statement
def test_customer_history_events_and_statement_pdf(wb, tmp_path):
    a = add(wb, **{"Customer Name": "شاہد محمود", "Mobile Number": "0300-1234567"})
    for m in ("2026-08", "2026-09", "2026-10"):
        gen(wb, m)
    pay(wb, a, 1300, month="2026-08", date=dt.date(2026, 8, 10))
    pay(wb, a, 500, month="2026-09", date=dt.date(2026, 9, 12))
    st = S_(wb)
    svc.update_customer(st, a, {"VLAN ID": "407"}, dt.date(2026, 10, 20))
    st.save()
    d = data(wb)
    h = svc.customer_history(d, a)
    assert [b["Month"] for b in h["bills"]] == ["2026-08", "2026-09", "2026-10"]
    assert (h["total_billing"], h["total_paid"], h["total_outstanding"]) == (3900, 1800, 2100)
    assert (h["last_payment_date"], h["last_payment_amount"]) == (dt.date(2026, 9, 12), 500)
    ev = [e for e in h["events"] if e["Event Type"] == "VLAN Change"][0]
    assert (ev["Old Value"], ev["New Value"]) == ("405", "407")
    hp = svc.customer_history(d, a, "2026-09", "2026-10")
    assert [b["Month"] for b in hp["bills"]] == ["2026-09", "2026-10"]
    pdf = pdfgen.statement_pdf(h, d.settings, tmp_path)
    assert pdf.name == "CL-0001_شاہد_محمود_Statement.pdf" and pdf.read_bytes()[:4] == b"%PDF"
    assert svc.search_customers(d, "cl-0001")[0]["Customer ID"] == a              # capitalisation ignored
    assert svc.search_customers(d, "03001234567")[0]["Customer ID"] == a
    assert svc.search_customers(d, "407", "VLAN ID")[0]["Customer ID"] == a
    assert "923001234567" in whatsapp_link(d.settings, "0300-1234567")


def test_customer_id_never_changes(wb):
    a = add(wb)
    st = S_(wb)
    with pytest.raises(UserError, match="can never be changed"):
        svc.update_customer(st, a, {"Customer ID": "CL-0100"})
    svc.update_customer(st, a, {"Location Code": "B2", "Package": "Premium", "Customer Status": "Suspended"})
    st.save()
    d = data(wb)
    assert d.customer(a)["Full Address"] == "Arai Colony, Gali 2"
    types = [e["Event Type"] for e in d.events]
    assert "Address Change" in types and "Package Change" in types and "Suspension" in types


def test_excel_edits_are_logged_as_events(wb):
    a = add(wb)
    st = S_(wb)
    c = st.rows(S.CUSTOMERS)[0]
    st.t(S.CUSTOMERS).ws.cell(c["_row"], 8).value = "499"        # VLAN edited directly in Excel
    st.save()
    st = S_(wb)
    found = svc.sync_events(st)
    st.save()
    assert found == [f"{a} VLAN ID: 405 -> 499"]
    assert data(wb).events[-1]["Event Type"] == "VLAN Change"


# --------------------------------------------------------------- safety
def test_backup_never_overwrites(wb, tmp_path):
    b1 = bk.backup(wb, tmp_path / "backups", "before_new_month")
    b2 = bk.backup(wb, tmp_path / "backups", "before_new_month")
    assert b1 != b2 and b1.exists() and b2.exists()
    assert b1.parent.name == dt.date.today().strftime("%Y-%m")


def test_app_write_makes_backup_and_refuses_open_workbook(wb):
    app = App(wb, root=wb.parent.parent)
    add(wb)
    app.write(lambda s: svc.generate_month(s, M1), reason="before_new_month_2026-10")
    assert any("before_new_month_2026-10" in p.name for p in bk.list_backups(app.backups))
    lock_file(wb).write_text("x")
    with pytest.raises(UserError, match="open in Excel"):
        app.write(lambda s: None)
    lock_file(wb).unlink()


def test_reports_export(wb):
    _area_setup(wb)
    app = App(wb, root=wb.parent.parent)
    out = export_reports(app, M1)
    from openpyxl import load_workbook
    names = load_workbook(out).sheetnames
    assert "Cash_Reconciliation" in names and "Monthly_Comparison" in names


# ---------------------------------------------------------- OCR / import
def test_handwritten_verification_queue(wb, tmp_path):
    txt = tmp_path / "page1.txt"
    txt.write_text("\n".join([
        "1 | Shahid Mahmood | A1 | 405 | 10 | 1300",
        "9 | ? | A1 | 406 | 10 | 1300",
        "12 | Bilal | A2 | 4?7 | 10 | 1500",
        "14 | Asif | B1 | 410 | 10 | ??",
        "15 | Nadeem | A7 | 411 | 10 | 1300",
    ]), encoding="utf-8")
    st = S_(wb)
    batch, n, k = imp.load_transcription(st, txt, "HW-TEST")
    st.save()
    assert (n, k) == (5, 4)
    st = S_(wb)
    q = {(int(i["Row No"]), i["Field"]) for i in imp.open_items(st)}
    assert q == {(9, "Customer Name"), (12, "VLAN ID"), (14, "Monthly Fee"), (15, "Location Code")}
    # nothing guessed: the staged values are exactly what was read
    staged = {int(r["Row No"]): r for r in imp.staged_rows(st, batch)}
    assert staged[12]["VLAN ID"] == "4?7" and staged[9]["Customer Name"] == "?"
    with pytest.raises(UserError, match="not accepted"):
        imp.apply_answers(st, imp.parse_answers("14 = abc"), batch)
    msgs = imp.apply_answers(st, imp.parse_answers("9 = خالد محمود\n12 = 407\n14 = 1500"), batch)
    st.save()
    assert len(msgs) == 3
    st = S_(wb)
    staged = {int(r["Row No"]): r for r in imp.staged_rows(st, batch)}
    assert staged[9]["Customer Name"] == "خالد محمود" and staged[12]["VLAN ID"] == "407" and staged[14]["Monthly Fee"] == "1500"
    # all fields valid is NOT enough: rows wait for the page review
    assert staged[1]["Row Status"] == "Awaiting Page Review" and staged[15]["Row Status"] == "Needs Verification"
    with pytest.raises(UserError, match="PAGE NOT REVIEWED"):
        imp.commit_batch(st, batch)
    with pytest.raises(UserError, match="open verification question"):       # row 15 still unclear
        imp.confirm_page_reviewed(st, batch, "Office", "PAGE REVIEWED")
    imp.reject_row(st, batch, 15, "location code unreadable - will re-check page")
    imp.confirm_page_reviewed(st, batch, "Office", "PAGE REVIEWED")
    st.save()
    st = S_(wb)
    assert {r["Row Status"] for r in imp.staged_rows(st, batch)} == {"Ready", "Rejected"}
    r = imp.commit_batch(st, batch)
    st.save()
    assert len(r["imported"]) == 4
    names = [c["Customer Name"] for c in data(wb).customers]
    assert "خالد محمود" in names and "Nadeem" not in names


def _clean_page(wb, tmp_path, batch="HW-PAGE"):
    txt = tmp_path / f"{batch}.txt"
    txt.write_text("1 | Shahid Mahmood | A1 | 405 | 10 | 1300\n2 | Imran | B1 | 207 | 10 | 1500", encoding="utf-8")
    st = S_(wb)
    imp.load_transcription(st, txt, batch)
    st.save()
    return batch


def test_page_review_required_even_when_all_fields_valid(wb, tmp_path):
    batch = _clean_page(wb, tmp_path)
    st = S_(wb)
    b = imp.get_batch(st, batch)
    assert b["Page Reviewed"] == "NO" and not imp.is_reviewed(b)
    assert {r["Row Status"] for r in imp.staged_rows(st, batch)} == {"Awaiting Page Review"}
    assert imp.open_items(st, batch) == []
    with pytest.raises(UserError, match="PAGE NOT REVIEWED"):
        imp.commit_batch(st, batch)
    assert data(wb).customers == []                                           # nothing auto-committed


def test_page_review_needs_exact_phrase_and_name(wb, tmp_path):
    batch = _clean_page(wb, tmp_path)
    st = S_(wb)
    for who, typed in (("Office", "yes"), ("Office", "reviewed"), ("Office", ""), ("", "PAGE REVIEWED")):
        with pytest.raises(UserError):
            imp.confirm_page_reviewed(st, batch, who, typed)
    assert not imp.is_reviewed(imp.get_batch(st, batch))
    imp.confirm_page_reviewed(st, batch, "Office", "page reviewed")            # capital letters do not matter
    st.save()
    st = S_(wb)
    b = imp.get_batch(st, batch)
    assert (b["Page Reviewed"], b["Reviewed By"]) == ("YES", "Office") and b["Reviewed At"]
    assert {r["Row Status"] for r in imp.staged_rows(st, batch)} == {"Ready"}
    r = imp.commit_batch(st, batch)
    st.save()
    assert [cid for _, cid in r["imported"]] == ["CL-0001", "CL-0002"]


def test_yes_typed_in_excel_alone_is_not_a_review(wb, tmp_path):
    batch = _clean_page(wb, tmp_path)
    st = S_(wb)
    b = imp.get_batch(st, batch)
    st.t(S.BATCHES).ws.cell(b["_row"], S.BATCHES.names.index("Page Reviewed") + 1).value = "YES"   # no name / time
    st.save()
    with pytest.raises(UserError, match="PAGE NOT REVIEWED"):
        imp.commit_batch(S_(wb), batch)


def test_changes_after_review_cancel_the_review(wb, tmp_path):
    batch = _clean_page(wb, tmp_path)
    st = S_(wb)
    imp.confirm_page_reviewed(st, batch, "Office", "PAGE REVIEWED")
    st.save()
    st = S_(wb)
    row = imp.staged_rows(st, batch)[0]
    st.t(S.STAGING).update(row["_row"], {"VLAN ID": "406"})                  # edited in Excel after the review
    st.save()
    st = S_(wb)
    r = imp.commit_batch(st, batch)
    st.save()
    assert r["review_cancelled"] and r["imported"] == []
    assert not imp.is_reviewed(imp.get_batch(S_(wb), batch))
    assert data(wb).customers == []


def test_old_workbook_is_upgraded_without_data_loss(wb):
    add(wb)
    st = S_(wb)
    del st.wb["IMPORT_BATCHES"]                                                # simulate a workbook made before this feature
    st.save()
    st = S_(wb)
    assert st.upgraded and "IMPORT_BATCHES" in st.wb.sheetnames
    st.save()
    assert [c["Customer ID"] for c in data(wb).customers] == ["CL-0001"]


def test_low_ocr_confidence_goes_to_queue(wb, tmp_path):
    from openpyxl import Workbook
    p = tmp_path / "ocr.xlsx"
    w = Workbook()
    ws = w.active
    ws.append(["Page", "Row No"] + imp.HW_FIELDS + ["Other Fields"] + [f"{f} Confidence" for f in imp.HW_FIELDS])
    ws.append([1, 3, "Shahid", "A1", "405", 10, 1300, "", "", 98, 99, 62, 97, 99, ""])
    w.save(p)
    st = S_(wb)
    _, _, k = imp.load_transcription(st, p, "HW-OCR")
    items = imp.open_items(st)
    assert k == 1 and items[0]["Field"] == "VLAN ID" and "62" in items[0]["Reason"]
    assert items[0]["Detected Value"] == "405"          # not changed, only queued


def test_answer_needs_field_when_row_has_several(wb, tmp_path):
    txt = tmp_path / "p.txt"
    txt.write_text("9 | ? | A1 | ? | 10 | 1300", encoding="utf-8")
    st = S_(wb)
    imp.load_transcription(st, txt, "HW-2")
    with pytest.raises(UserError, match="several open fields"):
        imp.apply_answers(st, imp.parse_answers("9 = Khalid"))
    imp.apply_answers(st, imp.parse_answers("9 name = Khalid\n9 vlan = 0407"))
    assert {r["VLAN ID"] for r in imp.staged_rows(st, "HW-2")} == {"0407"}


def test_import_existing_excel_preview(wb, tmp_path):
    add(wb, **{"Mobile Number": "0300-1111111"})
    f = tmp_path / "old.csv"
    f.write_text("Name,Location,VLAN,Due,Fee,Mobile\nNew Person,A2,300,10,1500,0300-2222222\n"
                 "Shahid Mahmood,A1,301,10,1300,\nBad Loc,A9,302,10,1300,\nNo Vlan,A1,,10,1300,\n"
                 "Mobile Dup,B1,303,5,1300,03001111111\n", encoding="utf-8")
    rows = imp.read_table_file(f)
    p = imp.preview_customers(Store(wb), rows)
    assert [c["Customer Name"] for c in p.new] == ["New Person"]
    assert {c["Customer Name"] for c, _ in p.duplicates} == {"Shahid Mahmood", "Mobile Dup"}
    assert [c["Customer Name"] for c, _ in p.invalid_location] == ["Bad Loc"]
    assert [c["Customer Name"] for c, _ in p.missing] == ["No Vlan"]
    st = S_(wb)
    ids = imp.commit_customers(st, p.new, "test import")
    st.save()
    assert ids == ["CL-0002"]


# ------------------------------------------- shared VLANs / location normalization / source marks
def test_shared_vlan_is_not_a_duplicate(wb):
    a = add(wb, **{"Customer Name": "Jawad", "Location Code": "A2", "VLAN ID": "405"})
    b = add(wb, **{"Customer Name": "Khalid", "Location Code": "A3", "VLAN ID": "405"})   # no confirmation needed
    c = add(wb, **{"Customer Name": "Riaz", "Location Code": "A1", "VLAN ID": "207"})
    d_ = add(wb, **{"Customer Name": "Aslam", "Location Code": "A2", "VLAN ID": "207"})
    d = data(wb)
    assert [x["VLAN ID"] for x in d.customers] == ["405", "405", "207", "207"]
    assert svc.find_possible_duplicates(d, {"Customer Name": "New", "Location Code": "B1", "VLAN ID": "405"}) == []
    assert svc.check_data(d) == []                                     # shared VLANs are not reported at all
    # stronger identity signals still work
    with pytest.raises(UserError, match="same Name \\+ Location Code"):
        add(wb, **{"Customer Name": "jawad", "Location Code": "A02", "VLAN ID": "999"})


def test_non_shared_vlan_reuse_is_info_only(wb):
    add(wb, **{"Customer Name": "P", "VLAN ID": "777"})
    add(wb, **{"Customer Name": "Q", "Location Code": "A2", "VLAN ID": "777"})          # still saved without a prompt
    issues = svc.check_data(data(wb))
    assert len(issues) == 1 and issues[0].startswith("Info: VLAN 777") and "Not a duplicate" in issues[0]


def test_location_code_equivalent_forms(wb):
    st = S_(wb)
    for n in range(4, 11):
        svc.add_location(st, f"A{n}", "")
    st.save()
    assert svc.normalize_location_code("A04") == "A4" and svc.normalize_location_code("a010") == "A10"
    cid = add(wb, **{"Location Code": "A04", "VLAN ID": "405"})
    c = data(wb).customer(cid)
    assert c["Location Code"] == "A4" and c["Area"] == "Arif Town"
    assert [b for b in svc.select_bills(data(wb), M1, location="A04")] == []      # no month yet, but no error


def test_handwritten_codes_kept_as_written_and_marks_not_in_notes(wb, tmp_path):
    st = S_(wb)
    svc.add_location(st, "A4", "")
    svc.add_location(st, "A10", "")
    st.save()
    txt = tmp_path / "p.txt"
    txt.write_text("1 | Ali | A04 | 405 | 10 | 1500\n2 | Bilal | A010 | 405 | 10 | 1300", encoding="utf-8")
    st = S_(wb)
    imp.load_transcription(st, txt, "HW-N")
    assert imp.open_items(st, "HW-N") == []                             # A04 / A010 accepted, 405 twice accepted
    rows = imp.staged_rows(st, "HW-N")
    assert [r["Location Code"] for r in rows] == ["A04", "A010"]        # original handwriting preserved
    srow = rows[0]
    st.t(S.STAGING).update(srow["_row"], {"Other Fields": "Mobile column: C-on"})
    imp.confirm_page_reviewed(st, "HW-N", "Office", "PAGE REVIEWED")
    r = imp.commit_batch(st, "HW-N")
    st.save()
    assert len(r["imported"]) == 2 and r["possible_duplicates"] == []   # shared VLAN 405: no duplicate prompt
    d = data(wb)
    assert [c["Location Code"] for c in d.customers] == ["A4", "A10"]
    assert all("C-on" not in (c.get("Notes") or "") for c in d.customers)
    assert "unresolved source marks kept in IMPORT_STAGING" in d.customers[0]["Notes"]
    assert imp.staged_rows(S_(wb), "HW-N")[0]["Other Fields"] == "Mobile column: C-on"


def test_old_workbook_gets_shared_vlan_setting(wb):
    add(wb)
    st = S_(wb)
    del st.wb.defined_names["CFG_SharedVLANs"]
    row = 5 + [x.key for x in S.SETTINGS].index("SharedVLANs")
    for c in range(1, 4):
        st.wb["SETTINGS"].cell(row, c).value = None
    st.save()
    st = S_(wb)
    assert any("Shared VLANs" in u for u in st.upgraded)
    st.save()
    assert svc.shared_vlans(S_(wb).settings()) == {"405", "207"}
    assert [c["Customer ID"] for c in data(wb).customers] == ["CL-0001"]


@pytest.mark.skipif(shutil.which("soffice") is None, reason="LibreOffice not installed")
def test_excel_data_check_ignores_shared_vlans(wb):
    sys.path.insert(0, str(ROOT / "tools"))
    import verify_workbook as vw
    from openpyxl import load_workbook
    for name, loc, vlan in (("A", "A1", "405"), ("B", "A2", "405"), ("C", "A3", "777"), ("D", "B1", "777")):
        add(wb, **{"Customer Name": name, "Location Code": loc, "VLAN ID": vlan})
    calc = vw.recalc(wb)
    ws = load_workbook(calc, data_only=True)["CUSTOMERS"]
    col = S.CUSTOMERS.names.index("Data Check") + 1
    checks = [ws.cell(r, col).value for r in range(5, 9)]
    assert checks[0] in (None, "") and checks[1] in (None, "")
    assert checks[2].startswith("NOTE: VLAN") and checks[3].startswith("NOTE: VLAN")


# ---------------------------------------------------------------- utilities
def test_months_and_filenames():
    assert parse_month("November 2026") == "2026-11" and parse_month("11/2026") == "2026-11" and parse_month("Nov-2026") == "2026-11"
    with pytest.raises(UserError):
        parse_month("2026-13")
    assert safe_filename('CL-0025 Shahid: "M"/x?') == "CL-0025_Shahid_Mx"
    assert safe_filename("CL-0001 شاہد محمود") == "CL-0001_شاہد_محمود"
    assert whatsapp_number("0300 1234567") == "923001234567" and whatsapp_number("12") is None


def test_demo_workbook_builds(tmp_path):
    p = build_demo(tmp_path / "demo.xlsx")
    d = svc.load(Store(p, check_lock=False))
    assert all(c["Customer Name"].startswith("[DEMO]") for c in d.customers)
    assert sorted({b["Month"] for b in d.billing}) == ["2026-08", "2026-09", "2026-10"]
    assert svc.check_data(d) == []


@pytest.mark.skipif(shutil.which("soffice") is None, reason="LibreOffice not installed")
def test_excel_formulas_match_python(tmp_path):
    sys.path.insert(0, str(ROOT / "tools"))
    import verify_workbook as vw
    p = build_demo(tmp_path / "demo.xlsx")
    assert vw.structure(p) == []
    calc = vw.recalc(p)
    assert calc is not None
    assert vw.compare(p, calc) == []
