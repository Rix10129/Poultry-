"""Importing customers.

1. Existing Excel/CSV lists  -> preview (new / possible duplicate / invalid location / missing fields)
                              -> import only after confirmation.
2. Handwritten recovery sheets (OCR or typed transcription)
       -> IMPORT_STAGING + VERIFICATION_QUEUE
       -> operator answers only the unclear fields ("9 = خالد محمود")
       -> commit verified rows as customers.

The system NEVER guesses unclear handwriting and never silently corrects a value.
"""
from __future__ import annotations

import csv
import datetime as dt
import re
from dataclasses import dataclass, field
from pathlib import Path

from openpyxl import Workbook, load_workbook

from . import schema as S
from . import services as svc
from .store import Store
from .util import UserError, clean_text, now, to_number

# ------------------------------------------------------------ column aliases
ALIASES = {
    "Customer ID": ["customer id", "id", "cust id", "customer no"],
    "Customer Name": ["customer name", "name", "customer", "naam", "نام"],
    "Location Code": ["location code", "location", "loc", "code", "area code", "location/street"],
    "VLAN ID": ["vlan id", "vlan", "vlan no"],
    "Due Date": ["due date", "due", "due day", "expiry", "expiry date", "date"],
    "Monthly Fee": ["monthly fee", "fee", "amount", "bill", "rent", "monthly bill"],
    "Fee Override": ["fee override", "special fee"],
    "Mobile Number": ["mobile number", "mobile", "phone", "cell", "contact", "mobile no"],
    "Package": ["package", "package name", "plan"],
    "Address Detail": ["address detail", "house", "house no"],
    "Connection Date": ["connection date", "connected on", "start date"],
    "Customer Status": ["customer status", "status"],
    "Notes": ["notes", "remarks", "note"],
    "Default Collector ID": ["default collector id", "collector id", "collector"],
}


def _match_header(h: str) -> str | None:
    k = clean_text(h).lower().replace("_", " ")
    for field_, names in ALIASES.items():
        if k in names:
            return field_
    return None


def read_table_file(path: Path) -> list[dict]:
    path = Path(path)
    if not path.exists():
        raise UserError(f"File not found: {path}")
    rows: list[list] = []
    if path.suffix.lower() in (".csv", ".txt"):
        with open(path, newline="", encoding="utf-8-sig") as f:
            rows = [r for r in csv.reader(f)]
    elif path.suffix.lower() in (".xlsx", ".xlsm"):
        wb = load_workbook(path, data_only=True, read_only=True)
        ws = wb.worksheets[0]
        rows = [list(r) for r in ws.iter_rows(values_only=True)]
    else:
        raise UserError("Only .xlsx and .csv files can be imported.")
    # find header row = first row with >= 2 recognised headers
    for hi, r in enumerate(rows[:20]):
        mapped = [_match_header(clean_text(x)) if x is not None else None for x in r]
        if sum(1 for m in mapped if m) >= 2:
            break
    else:
        raise UserError("Could not find a header row. The first row should contain column names like "
                        "Customer Name, Location Code, VLAN ID, Due Date, Monthly Fee.")
    out = []
    for n, r in enumerate(rows[hi + 1:], start=hi + 2):
        if not any(clean_text(x) for x in r):
            continue
        d = {"_source_row": n}
        for m, v in zip(mapped, r):
            if m:
                d[m] = v
        out.append(d)
    return out


@dataclass
class Preview:
    new: list[dict] = field(default_factory=list)
    duplicates: list[tuple[dict, list]] = field(default_factory=list)
    invalid_location: list[tuple[dict, str]] = field(default_factory=list)
    missing: list[tuple[dict, list]] = field(default_factory=list)


def prepare_customer(r: dict) -> dict:
    c = {k: v for k, v in r.items() if not k.startswith("_")}
    for k in ("Customer ID", "Customer Name", "Location Code", "VLAN ID", "Mobile Number", "Package", "Customer Status"):
        if k in c:
            c[k] = clean_text(c[k])
    if c.get("Location Code"):
        c["Location Code"] = c["Location Code"].upper().replace(" ", "")
    if "Monthly Fee" in c and c.get("Fee Override") in (None, ""):
        c["Fee Override"] = c.pop("Monthly Fee")
    else:
        c.pop("Monthly Fee", None)
    c["Customer Status"] = (c.get("Customer Status") or "Active").title() if c.get("Customer Status") else "Active"
    if c["Customer Status"] == "On hold":
        c["Customer Status"] = "On Hold"
    return c


def preview_customers(store: Store, rows: list[dict]) -> Preview:
    d = svc.load(store)
    p = Preview()
    batch_seen: list[dict] = []
    locs = svc.location_map(d)
    for r in rows:
        c = prepare_customer(r)
        c["_source_row"] = r.get("_source_row")
        missing = [f for f in ("Customer Name", "Location Code", "VLAN ID", "Due Date") if c.get(f) in (None, "")]
        if to_number(c.get("Fee Override")) is None and not c.get("Package"):
            missing.append("Monthly Fee")
        if missing:
            p.missing.append((c, missing))
            continue
        if c["Location Code"] not in locs:
            p.invalid_location.append((c, f"Location Code {c['Location Code']} does not exist."))
            continue
        errs = svc.validate_customer(d, c)
        if errs:
            p.missing.append((c, errs))
            continue
        dups = svc.find_possible_duplicates(d, c)
        for o in batch_seen:
            if svc.norm_key(o["Customer Name"]) == svc.norm_key(c["Customer Name"]) and o["Location Code"] == c["Location Code"]:
                dups.append((f"file row {o['_source_row']}", "same Name + Location Code in this file"))
            elif c.get("Mobile Number") and svc.digits(o.get("Mobile Number")) == svc.digits(c["Mobile Number"]):
                dups.append((f"file row {o['_source_row']}", "same Mobile Number in this file"))
        batch_seen.append(c)
        if dups:
            p.duplicates.append((c, dups))
        else:
            p.new.append(c)
    return p


def write_preview_report(p: Preview, path: Path) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "SUMMARY"
    ws.append(["Category", "Rows"])
    for k, v in (("New customers", len(p.new)), ("Possible duplicates", len(p.duplicates)),
                 ("Invalid location codes", len(p.invalid_location)), ("Missing required fields", len(p.missing))):
        ws.append([k, v])
    cols = ["_source_row", "Customer ID", "Customer Name", "Location Code", "VLAN ID", "Due Date", "Fee Override", "Package", "Mobile Number"]
    for title, items in (("NEW", [(c, "") for c in p.new]),
                         ("POSSIBLE_DUPLICATES", [(c, "; ".join(f"{i}: {r}" for i, r in d)) for c, d in p.duplicates]),
                         ("INVALID_LOCATION", p.invalid_location),
                         ("MISSING_FIELDS", [(c, ", ".join(m)) for c, m in p.missing])):
        s = wb.create_sheet(title)
        s.append(["File Row"] + cols[1:] + ["Problem"])
        for c, why in items:
            s.append([clean_text(c.get(k)) for k in cols] + [why])
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


def commit_customers(store: Store, customers: list[dict], note: str) -> list[str]:
    ids = []
    for c in customers:
        c = {k: v for k, v in c.items() if not k.startswith("_")}
        ids.append(svc.add_customer(store, c, confirm_duplicates=lambda _d: True, event_note=note))
    return ids


# =============================================== handwritten / OCR import
HW_FIELDS = ["Customer Name", "Location Code", "VLAN ID", "Due Date", "Monthly Fee", "Mobile Number"]
REQUIRED_HW = ["Customer Name", "Location Code", "VLAN ID", "Due Date", "Monthly Fee"]
UNCLEAR_MARKS = re.compile(r"^\s*(\?+|unclear|illegible|\[unclear\]|xx+|n/?a)\s*$|\?", re.I)
FIELD_SHORT = {"name": "Customer Name", "customer name": "Customer Name", "loc": "Location Code", "location": "Location Code",
               "location code": "Location Code", "vlan": "VLAN ID", "vlan id": "VLAN ID", "due": "Due Date",
               "due date": "Due Date", "fee": "Monthly Fee", "monthly fee": "Monthly Fee", "mobile": "Mobile Number",
               "mobile number": "Mobile Number"}


def create_transcription_template(path: Path) -> Path:
    """Blank sheet for typing (or OCR output of) one handwritten page."""
    wb = Workbook()
    ws = wb.active
    ws.title = "TRANSCRIPTION"
    ws.append(["Page", "Row No"] + HW_FIELDS + ["Other Fields"] + [f"{f} Confidence" for f in HW_FIELDS])
    ws.append([1, 1, "", "", "", "", "", "", "", "", "", "", "", "", ""])
    info = wb.create_sheet("HOW_TO_FILL")
    for line in [
        "Copy each handwritten row exactly as written. One row per customer.",
        "If ANY value is unclear, type ? (or leave it blank). NEVER guess.",
        "Confidence columns are optional (0-100). OCR programs fill them; a person typing can leave them blank.",
        "Location Code examples: A1 = Arif Town Gali 1, B3 = Arai Colony Gali 3.",
        "Due Date = day of month (e.g. 10). Monthly Fee = number only (e.g. 1300).",
        "Customer names may be in Urdu or English - type them exactly as written.",
    ]:
        info.append([line])
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


def parse_handwritten_line(text: str) -> tuple[dict, list[str]]:
    """Split a typed/OCR line like '1 | Shahid Mahmood | A1 | 405 | 10 | 1300' or
    'Shahid Mahmood A1 405 10 1300'. Returns (fields, problems). Nothing is guessed:
    if the structure is not exactly recognisable, the fields are left unclear."""
    s = clean_text(text)
    out = {f: "" for f in HW_FIELDS}
    problems = []
    if "|" in s or "\t" in s or "," in s:
        parts = [clean_text(x) for x in re.split(r"[|\t,]", s)]
        if parts and re.fullmatch(r"\d{1,3}", parts[0]) and len(parts) >= 6:
            out["_row"] = int(parts[0])
            parts = parts[1:]
        for f, v in zip(["Customer Name", "Location Code", "VLAN ID", "Due Date", "Monthly Fee", "Mobile Number"], parts):
            out[f] = v
        if len(parts) < 5:
            problems.append("line has fewer than 5 columns")
        return out, problems
    toks = s.split()
    if toks and re.fullmatch(r"\d{1,3}", toks[0]) and len(toks) >= 6:
        out["_row"] = int(toks[0])
        toks = toks[1:]
    loc_idx = [i for i, t in enumerate(toks) if svc.LOC_RE.match(t) and re.match(r"^[A-Za-z]{1,3}\d{1,3}$", t)]
    if len(loc_idx) != 1:
        problems.append("could not find exactly one Location Code")
        out["Customer Name"] = s
        return out, problems
    i = loc_idx[0]
    out["Customer Name"] = " ".join(toks[:i])
    out["Location Code"] = toks[i]
    nums = toks[i + 1:]
    for f, v in zip(["VLAN ID", "Due Date", "Monthly Fee", "Mobile Number"], nums):
        out[f] = v
    if len(nums) < 3:
        problems.append("expected VLAN, Due Date and Fee after the Location Code")
    return out, problems


def _field_problem(fieldname: str, value: str, conf, threshold: float, known_locs: set) -> str | None:
    v = clean_text(value)
    if not v:
        return "missing / not readable" if fieldname in REQUIRED_HW else None
    if UNCLEAR_MARKS.search(v):
        return "marked unclear"
    if conf not in (None, "") and to_number(conf) is not None and to_number(conf) < threshold:
        return f"low OCR confidence ({to_number(conf)}%)"
    if fieldname == "Location Code":
        code = v.upper().replace(" ", "")
        if not svc.LOC_RE.match(code):
            return "not a valid location code format"
        if code not in known_locs:
            return f"location code {code} does not exist"
    if fieldname == "VLAN ID" and not re.fullmatch(r"[0-9A-Za-z\-/]+", v):
        return "VLAN contains unexpected characters"
    if fieldname == "Due Date":
        n = to_number(v)
        if n is None or not float(n).is_integer() or not 1 <= n <= 31:
            return "due date must be 1-31"
    if fieldname == "Monthly Fee":
        n = to_number(v)
        if n is None or n <= 0:
            return "fee is not a valid amount"
    if fieldname == "Mobile Number" and not re.fullmatch(r"[0-9+\- ]{7,16}", v):
        return "mobile number format looks wrong"
    return None


def load_transcription(store: Store, path: Path, batch_id: str | None = None, threshold: float | None = None,
                       source_label: str | None = None) -> tuple[str, int, int]:
    """Stage a transcription (xlsx/csv with the template columns, or .txt with one line per row).
    Returns (batch_id, rows, open_items)."""
    path = Path(path)
    d = svc.load(store)
    thr = threshold if threshold is not None else (to_number(d.settings.get("OCRConfidenceThreshold")) or 90)
    known = set(svc.location_map(d))
    batch_id = batch_id or f"HW-{dt.datetime.now():%Y%m%d-%H%M%S}"
    if any(r["Batch ID"] == batch_id for r in store.rows(S.STAGING)):
        raise UserError(f"Batch {batch_id} already exists.")
    entries = []          # (page, row_no, values, confidences, other, extra_problems)
    if path.suffix.lower() == ".txt":
        for n, line in enumerate(path.read_text(encoding="utf-8-sig").splitlines(), start=1):
            if not clean_text(line):
                continue
            vals, probs = parse_handwritten_line(line)
            entries.append((1, vals.pop("_row", n), vals, {}, "", probs))
    else:
        if path.suffix.lower() == ".csv":
            with open(path, newline="", encoding="utf-8-sig") as f:
                raw = list(csv.reader(f))
        else:
            raw = [list(r) for r in load_workbook(path, data_only=True).worksheets[0].iter_rows(values_only=True)]
        header = [clean_text(h) for h in raw[0]]
        for n, r in enumerate(raw[1:], start=1):
            row = dict(zip(header, r))
            vals = {f: clean_text(row.get(f)) for f in HW_FIELDS}
            if not any(vals.values()):
                continue
            conf = {f: row.get(f"{f} Confidence") for f in HW_FIELDS}
            entries.append((clean_text(row.get("Page")) or "1", int(to_number(row.get("Row No")) or n), vals, conf,
                            clean_text(row.get("Other Fields")), []))
    if not entries:
        raise UserError("No rows found in the file.")
    st, vq = store.t(S.STAGING), store.t(S.VERIFY)
    existing = store.rows(S.VERIFY)
    k = int(svc._next_seq(existing, "Item ID", "VQ", 5).split("-")[1])
    open_items = 0
    for page, row_no, vals, conf, other, probs in entries:
        items = []
        for f in HW_FIELDS:
            why = _field_problem(f, vals.get(f, ""), conf.get(f), thr, known)
            if why:
                items.append((f, why))
        if probs and not items:
            items.append(("Customer Name", "; ".join(probs)))
        status = "Needs Verification" if items else "Ready"
        st.append({"Batch ID": batch_id, "Row No": row_no, "Page": page, **{f: vals.get(f, "") for f in HW_FIELDS},
                   "Other Fields": other, "Source File": source_label or path.name, "Row Status": status,
                   "Notes": "; ".join(probs)})
        for f, why in items:
            vq.append({"Item ID": f"VQ-{k:05d}", "Batch ID": batch_id, "Row No": row_no, "Field": f,
                       "Detected Value": vals.get(f, "") or "(blank)",
                       "Confidence": clean_text(conf.get(f)) if conf.get(f) not in (None, "") else "",
                       "Reason": why, "Question": f"Please enter {f} for Row {row_no}.", "Status": "Open"})
            k += 1
            open_items += 1
    return batch_id, len(entries), open_items


def open_items(store: Store, batch_id: str | None = None) -> list[dict]:
    return [v for v in store.rows(S.VERIFY) if v["Status"] == "Open" and (not batch_id or v["Batch ID"] == batch_id)]


def parse_answers(text: str) -> list[tuple[int, str | None, str]]:
    """'9 = خالد محمود' / '12 vlan = 407' / '14.fee=1500' -> [(row, field|None, value)]"""
    out = []
    for line in text.splitlines():
        line = line.strip()
        if not line:
            continue
        m = re.match(r"^\s*(\d+)\s*(?:[.\-:]?\s*([A-Za-z][A-Za-z ]*?))?\s*=\s*(.+?)\s*$", line)
        if not m:
            raise UserError(f"Could not understand '{line}'. Use: 9 = value   or   9 vlan = value")
        fld = None
        if m.group(2):
            fld = FIELD_SHORT.get(m.group(2).strip().lower())
            if not fld:
                raise UserError(f"Unknown field '{m.group(2)}' in '{line}'. Use name, loc, vlan, due, fee or mobile.")
        out.append((int(m.group(1)), fld, m.group(3)))
    return out


def apply_answers(store: Store, answers: list[tuple[int, str | None, str]], batch_id: str | None = None) -> list[str]:
    """Update ONLY the asked fields. Each answer is validated; invalid answers are refused."""
    d = svc.load(store)
    known = set(svc.location_map(d))
    items = open_items(store, batch_id)
    msgs = []
    for row_no, fld, value in answers:
        cand = [i for i in items if int(to_number(i["Row No"])) == row_no and (fld is None or i["Field"] == fld)]
        if not batch_id:
            batches = {i["Batch ID"] for i in cand}
            if len(batches) > 1:
                raise UserError(f"Row {row_no} is open in several batches ({', '.join(sorted(batches))}). Choose the batch first.")
        if not cand:
            raise UserError(f"There is no open question for Row {row_no}" + (f" / {fld}" if fld else "") + ".")
        if len(cand) > 1:
            raise UserError(f"Row {row_no} has several open fields ({', '.join(i['Field'] for i in cand)}). "
                            f"Answer like: {row_no} vlan = 407")
        item = cand[0]
        why = _field_problem(item["Field"], value, None, 0, known)
        if why:
            raise UserError(f"Row {row_no} {item['Field']}: '{value}' was not accepted ({why}). Nothing was changed for this row.")
        val = clean_text(value)
        if item["Field"] == "Location Code":
            val = val.upper().replace(" ", "")
        _set_staging_field(store, item["Batch ID"], row_no, item["Field"], val)
        store.t(S.VERIFY).update(item["_row"], {"Answer": val, "Status": "Resolved", "Resolved At": now()})
        items.remove(item)
        msgs.append(f"Row {row_no} {item['Field']} = {val}")
    _refresh_row_status(store)
    return msgs


def apply_sheet_answers(store: Store) -> list[str]:
    """Answers typed directly into VERIFICATION_QUEUE > Answer."""
    answers = [(int(to_number(i["Row No"])), i["Field"], clean_text(i["Answer"]), i["Batch ID"])
               for i in open_items(store) if clean_text(i.get("Answer"))]
    msgs = []
    for row, fld, val, batch in answers:
        msgs += apply_answers(store, [(row, fld, val)], batch)
    return msgs


def _set_staging_field(store, batch, row_no, fieldname, value):
    for r in store.rows(S.STAGING):
        if r["Batch ID"] == batch and int(to_number(r["Row No"])) == row_no:
            old = clean_text(r.get(fieldname))
            store.t(S.STAGING).update(r["_row"], {fieldname: value, "Notes": (clean_text(r.get("Notes")) + f" | {fieldname}: '{old}' verified as '{value}'").strip(" |")})
            return
    raise UserError(f"Row {row_no} not found in batch {batch}.")


def _refresh_row_status(store):
    open_rows = {(i["Batch ID"], int(to_number(i["Row No"]))) for i in open_items(store)}
    for r in store.rows(S.STAGING):
        if r["Row Status"] in ("Imported", "Rejected"):
            continue
        new = "Needs Verification" if (r["Batch ID"], int(to_number(r["Row No"]))) in open_rows else (
            "Possible Duplicate" if r["Row Status"] == "Possible Duplicate" else "Ready")
        if new != r["Row Status"]:
            store.t(S.STAGING).update(r["_row"], {"Row Status": new})


def staged_rows(store: Store, batch_id: str) -> list[dict]:
    return [r for r in store.rows(S.STAGING) if r["Batch ID"] == batch_id]


def commit_batch(store: Store, batch_id: str, confirm_duplicate=None, default_status="Active") -> dict:
    """Create customers for Ready rows. Rows that look like existing customers need
    ``confirm_duplicate(row, dups) -> True`` or they stay as 'Possible Duplicate'."""
    rows = staged_rows(store, batch_id)
    if not rows:
        raise UserError(f"Batch {batch_id} was not found.")
    res = {"imported": [], "waiting_verification": 0, "possible_duplicates": [], "errors": []}
    for r in rows:
        if r["Row Status"] in ("Imported", "Rejected"):
            continue
        if r["Row Status"] == "Needs Verification":
            res["waiting_verification"] += 1
            continue
        d = svc.load(store)
        c = {"Customer Name": r["Customer Name"], "Location Code": r["Location Code"].upper(), "VLAN ID": r["VLAN ID"],
             "Due Date": to_number(r["Due Date"]), "Fee Override": to_number(r["Monthly Fee"]),
             "Mobile Number": r["Mobile Number"], "Customer Status": default_status,
             "Notes": f"Imported from handwritten sheet {batch_id} row {r['Row No']}" + (f"; other: {r['Other Fields']}" if clean_text(r.get('Other Fields')) else "")}
        errs = svc.validate_customer(d, c)
        if errs:
            res["errors"].append(f"Row {r['Row No']}: " + "; ".join(errs))
            continue
        dups = svc.find_possible_duplicates(d, c)
        if dups and not (confirm_duplicate and confirm_duplicate(r, dups)):
            store.t(S.STAGING).update(r["_row"], {"Row Status": "Possible Duplicate",
                                                  "Notes": "; ".join(f"{i}: {w}" for i, w in dups)})
            res["possible_duplicates"].append((r["Row No"], dups))
            continue
        cid = svc.add_customer(store, c, confirm_duplicates=lambda _d: True,
                               event_note=f"Imported from handwritten sheet {batch_id}")
        store.t(S.STAGING).update(r["_row"], {"Row Status": "Imported", "Customer ID Assigned": cid})
        res["imported"].append((r["Row No"], cid))
    return res


def reject_row(store: Store, batch_id: str, row_no: int, reason: str):
    for r in staged_rows(store, batch_id):
        if int(to_number(r["Row No"])) == row_no:
            store.t(S.STAGING).update(r["_row"], {"Row Status": "Rejected", "Notes": reason})
            for i in open_items(store, batch_id):
                if int(to_number(i["Row No"])) == row_no:
                    store.t(S.VERIFY).update(i["_row"], {"Status": "Rejected", "Resolved At": now()})
            return
    raise UserError(f"Row {row_no} not found in batch {batch_id}.")


# --------------------------------------------------------------- OCR engine
def ocr_image(image_path: Path, threshold: float = 90) -> Path:
    """Optional: run Tesseract OCR (if installed with the Urdu 'urd' language) on a scanned sheet
    and write a transcription file with confidences. Every value below the threshold, and every
    line that cannot be split exactly, goes to the Verification Queue when loaded.

    Handwriting OCR is unreliable; expect most names to need verification."""
    try:
        import pytesseract
        from PIL import Image
    except ImportError:
        raise UserError("OCR is not installed (pytesseract + Tesseract with Urdu). Use the transcription template instead: "
                        "type the page into the template (write ? for anything unclear) and load it.")
    try:
        data = pytesseract.image_to_data(Image.open(image_path), lang="urd+eng", output_type=pytesseract.Output.DICT)
    except Exception as e:  # tesseract binary / language missing
        raise UserError(f"OCR failed: {e}")
    lines: dict = {}
    for i, word in enumerate(data["text"]):
        if not clean_text(word):
            continue
        key = (data["block_num"][i], data["par_num"][i], data["line_num"][i])
        lines.setdefault(key, []).append((word, float(data["conf"][i])))
    out = Path(image_path).with_suffix(".ocr.xlsx")
    wb = Workbook()
    ws = wb.active
    ws.append(["Page", "Row No"] + HW_FIELDS + ["Other Fields"] + [f"{f} Confidence" for f in HW_FIELDS])
    for n, (_, words) in enumerate(sorted(lines.items()), start=1):
        text = " ".join(w for w, _ in words)
        conf = min(c for _, c in words)
        vals, probs = parse_handwritten_line(text)
        ws.append([1, vals.pop("_row", n)] + [vals.get(f, "") for f in HW_FIELDS] + ["; ".join(probs)] + [conf] * len(HW_FIELDS))
    wb.save(out)
    return out
