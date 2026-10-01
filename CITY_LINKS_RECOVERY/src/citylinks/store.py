"""Workbook storage layer.

Reads/writes the Excel Tables defined in :mod:`schema`.  Python never relies on
cached formula results (openpyxl cannot calculate); it only reads *input* and
*system* columns and calculates everything else itself (see services.py).
"""
from __future__ import annotations

import datetime as dt
import os
import re
import shutil
import tempfile
import warnings
import zipfile
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Alignment, PatternFill
from openpyxl.utils import get_column_letter, range_boundaries

from . import schema as S
from .util import UserError, clean_text

FILL_CALC = PatternFill("solid", fgColor="F2F2F2")
FILL_SYSTEM = PatternFill("solid", fgColor="EAF1FB")

# Sheets whose print area grows/shrinks with the data: sheet -> (cell holding row count, last column)
DYNAMIC_PRINT_AREAS = {
    "PRINT_SHEET": ("$Z$1", 12),
    "PRINT_SHEET_UR": ("$Z$1", 12),
    "STATEMENT": ("$Z$1", 8),
}


class WorkbookOpenError(UserError):
    pass


def lock_file(path: Path) -> Path:
    return path.with_name("~$" + path.name)


class TableIO:
    def __init__(self, store: "Store", tdef: S.TableDef):
        self.store, self.tdef = store, tdef
        self.ws = store.wb[tdef.sheet]
        if tdef.name not in self.ws.tables:
            raise UserError(f"Table {tdef.name} is missing from sheet {tdef.sheet}. The workbook may be damaged.")
        self.table = self.ws.tables[tdef.name]
        self._read_bounds()

    def _read_bounds(self):
        c1, r1, c2, r2 = range_boundaries(self.table.ref)
        self.c1, self.header_row, self.c2, self.last_row = c1, r1, c2, r2
        self.colmap = {}
        for c in range(c1, c2 + 1):
            name = clean_text(self.ws.cell(r1, c).value)
            self.colmap[name] = c
        missing = [n for n in self.tdef.names if n not in self.colmap]
        if missing:
            raise UserError(f"Sheet {self.tdef.sheet} is missing column(s): {', '.join(missing)}. "
                            "Do not rename or delete header cells.")

    # ------------------------------------------------------------------ read
    def _row_is_blank(self, r: int) -> bool:
        for col in self.tdef.stored():
            v = self.ws.cell(r, self.colmap[col.name]).value
            if v not in (None, ""):
                return False
        return True

    def rows(self) -> list[dict]:
        out = []
        for r in range(self.header_row + 1, self.last_row + 1):
            if self._row_is_blank(r):
                continue
            d = {"_row": r}
            for col in self.tdef.stored():
                v = self.ws.cell(r, self.colmap[col.name]).value
                if isinstance(v, str) and v.startswith("="):
                    v = None          # never trust formulas typed into input columns
                if col.text:
                    v = clean_text(v)
                elif isinstance(v, str):
                    v = v.strip()
                d[col.name] = v
            out.append(d)
        return out

    # ----------------------------------------------------------------- write
    def _write_cell(self, r: int, col: S.Col, value):
        cell = self.ws.cell(r, self.colmap[col.name])
        if col.kind == S.CALC:
            cell.value = self.tdef.expand(col.formula)
            cell.fill = FILL_CALC
        else:
            if col.text and value not in (None, ""):
                value = clean_text(value)
                cell.number_format = S.FMT_TEXT
            elif isinstance(value, dt.datetime) and col.fmt == S.FMT_DATE:
                value = value.date()
            cell.value = None if value == "" else value
            if col.kind == S.SYSTEM:
                cell.fill = FILL_SYSTEM
        if col.fmt and not col.text:
            cell.number_format = col.fmt
        cell.alignment = Alignment(vertical="center", horizontal="left" if col.text else None)

    def append(self, values: dict) -> int:
        unknown = set(values) - set(self.tdef.names)
        if unknown:
            raise ValueError(f"Unknown columns for {self.tdef.name}: {unknown}")
        # Reuse a trailing blank row (a new table always has one), else grow the table.
        r = self.last_row
        if not self._row_is_blank(r):
            r = self.last_row + 1
            for c in range(self.c1, self.c2 + 1):
                if self.ws.cell(r, c).value not in (None, ""):
                    raise UserError(f"Cannot add a row to {self.tdef.sheet}: row {r} below the table is not empty.")
            self.last_row = r
            ref = f"{get_column_letter(self.c1)}{self.header_row}:{get_column_letter(self.c2)}{r}"
            self.table.ref = ref
            if self.table.autoFilter is not None:
                self.table.autoFilter.ref = ref
        for col in self.tdef.columns:
            self._write_cell(r, col, values.get(col.name))
        self.ws.row_dimensions[r].height = None
        return r

    def update(self, r: int, values: dict):
        for name, v in values.items():
            col = self.tdef.col(name)
            if col.kind == S.CALC:
                raise ValueError(f"{name} is a calculated column")
            self._write_cell(r, col, v)

    def ensure_formulas(self):
        """Re-write calculated column formulas on every row (repairs accidental overwrites)."""
        for r in range(self.header_row + 1, self.last_row + 1):
            for col in self.tdef.columns:
                if col.kind == S.CALC:
                    self._write_cell(r, col, None)


class Store:
    def __init__(self, path, wb=None, check_lock=True):
        self.path = Path(path)
        if wb is None:
            if not self.path.exists():
                raise UserError(f"Workbook not found: {self.path}")
            if check_lock and lock_file(self.path).exists():
                raise WorkbookOpenError(
                    f"{self.path.name} is open in Excel. Please SAVE and CLOSE it in Excel, then try again.")
            with warnings.catch_warnings():
                warnings.simplefilter("ignore", UserWarning)    # formula print areas are re-added on save
                wb = load_workbook(self.path)
        self.wb = wb
        self._tables: dict[str, TableIO] = {}
        self.upgraded: list[str] = []
        if "HOME" in wb.sheetnames and ("IMPORT_BATCHES" not in wb.sheetnames or "CFG_SharedVLANs" not in wb.defined_names):
            from .builder import upgrade_workbook          # older workbook: add new sheets, keep all data
            self.upgraded = upgrade_workbook(self)
        if "HOME" in wb.sheetnames:
            from .builder import sync_calc_formulas         # formulas changed in a newer version
            self.upgraded += sync_calc_formulas(self)

    def t(self, tdef: S.TableDef) -> TableIO:
        if tdef.name not in self._tables:
            self._tables[tdef.name] = TableIO(self, tdef)
        return self._tables[tdef.name]

    def rows(self, tdef: S.TableDef) -> list[dict]:
        return self.t(tdef).rows()

    # -------------------------------------------------------------- settings
    def _setting_cell(self, key: str):
        dn = self.wb.defined_names.get(f"CFG_{key}")
        if dn is None:
            raise UserError(f"Setting {key} is missing from SETTINGS.")
        sheet, ref = list(dn.destinations)[0]
        return self.wb[sheet][ref.replace("$", "")]

    def setting(self, key: str, default=None):
        try:
            v = self._setting_cell(key).value
        except UserError:
            return default if default is not None else S.SETTING_KEYS[key].default
        if v in (None, ""):
            return default if default is not None else ("" if key == "CurrentMonth" else S.SETTING_KEYS[key].default)
        return v

    def settings(self) -> dict:
        return {s.key: self.setting(s.key) for s in S.SETTINGS}

    def set_setting(self, key: str, value):
        self._setting_cell(key).value = value

    def named_cell(self, name: str):
        dn = self.wb.defined_names.get(name)
        sheet, ref = list(dn.destinations)[0]
        return self.wb[sheet][ref.replace("$", "").split(":")[0]]

    # ------------------------------------------------------------------ save
    def save(self, path=None):
        target = Path(path or self.path)
        if lock_file(target).exists():
            raise WorkbookOpenError(
                f"{target.name} is open in Excel. Please SAVE and CLOSE it in Excel, then try again. Nothing was saved.")
        target.parent.mkdir(parents=True, exist_ok=True)
        self.wb.calculation.fullCalcOnLoad = True
        fd, tmp = tempfile.mkstemp(suffix=".xlsx", dir=target.parent)
        os.close(fd)
        try:
            self.wb.save(tmp)
            inject_dynamic_print_areas(tmp)
            try:
                os.replace(tmp, target)
            except PermissionError:
                raise WorkbookOpenError(
                    f"Could not save {target.name} - it is probably open in Excel. Close it and try again. Nothing was saved.")
        finally:
            if os.path.exists(tmp):
                os.remove(tmp)


def inject_dynamic_print_areas(xlsx_path):
    """openpyxl cannot keep formula-based print areas, so they are re-added after every save.

    Excel then prints only the rows that contain data (plus totals/signatures)."""
    with zipfile.ZipFile(xlsx_path) as z:
        items = [(i, z.read(i.filename)) for i in z.infolist()]
    data = dict((i.filename, d) for i, d in items)
    xml = data["xl/workbook.xml"].decode("utf-8")
    sheets = re.findall(r'<sheet [^>]*name="([^"]+)"', xml)
    names = []
    for sheet, (cell, ncols) in DYNAMIC_PRINT_AREAS.items():
        if sheet not in sheets:
            continue
        idx = sheets.index(sheet)
        xml = re.sub(rf'<definedName name="_xlnm.Print_Area" localSheetId="{idx}"[^>]*>[^<]*</definedName>', "", xml)
        q = f"'{sheet}'"
        names.append(f'<definedName name="_xlnm.Print_Area" localSheetId="{idx}">'
                     f'OFFSET({q}!$A$1,0,0,MAX(1,{q}!{cell}),{ncols})</definedName>')
    if not names:
        return
    if "<definedNames>" in xml:
        xml = xml.replace("<definedNames>", "<definedNames>" + "".join(names), 1)
    elif "<definedNames/>" in xml:
        xml = xml.replace("<definedNames/>", "<definedNames>" + "".join(names) + "</definedNames>", 1)
    else:
        xml = xml.replace("</sheets>", "</sheets><definedNames>" + "".join(names) + "</definedNames>", 1)
    data["xl/workbook.xml"] = xml.encode("utf-8")
    tmp = str(xlsx_path) + ".tmp"
    with zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as z:
        for info, _ in items:
            z.writestr(info, data[info.filename])
    shutil.move(tmp, xlsx_path)
