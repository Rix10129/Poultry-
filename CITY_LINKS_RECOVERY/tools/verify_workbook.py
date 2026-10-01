"""Verify a CITY LINKS workbook.

1. Structure: sheets, tables, columns, formulas, named cells, dropdowns, print setup, protection.
2. Recalculation (needs LibreOffice 'soffice'): recalculates every formula and compares the
   Excel results with the Python business logic (billing, dashboard, cash, print list, history).

    python tools/verify_workbook.py data/DEMO_CITY_LINKS_Recovery.xlsx
"""
from __future__ import annotations

import os
import shutil
import subprocess
import sys
import tempfile
import warnings
import zipfile
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "src"))
warnings.filterwarnings("ignore")

from openpyxl import load_workbook  # noqa: E402

from citylinks import schema as S  # noqa: E402
from citylinks import services as svc  # noqa: E402
from citylinks.builder import SHEET_ORDER  # noqa: E402
from citylinks.store import Store  # noqa: E402


def structure(path: Path) -> list[str]:
    errs = []
    wb = load_workbook(path)
    for s in SHEET_ORDER:
        if s not in wb.sheetnames:
            errs.append(f"missing sheet {s}")
    for t in S.ALL_TABLES:
        ws = wb[t.sheet]
        if t.name not in ws.tables:
            errs.append(f"missing table {t.name}")
            continue
        tbl = ws.tables[t.name]
        hdr = [c.value for c in ws[4]]
        for col in t.columns:
            if col.name not in hdr:
                errs.append(f"{t.name}: missing column {col.name}")
        tc = {c.name: c for c in tbl.tableColumns}
        names = [c.name for c in tbl.tableColumns]
        if names != [h for h in hdr if h is not None][:len(names)]:
            errs.append(f"{t.name}: table column names do not match the header row (Excel would report damage)")
        for col in t.columns:
            if col.kind == S.CALC and (tc.get(col.name) is None or tc[col.name].calculatedColumnFormula is None):
                errs.append(f"{t.name}[{col.name}]: no calculated column formula")
        if not ws.data_validations.dataValidation and any(c.dv for c in t.columns):
            errs.append(f"{t.sheet}: no data validation")
        if not ws.print_title_rows:
            errs.append(f"{t.sheet}: no repeated header row for printing")
    for nm in ["CFG_CompanyName", "CFG_CurrentMonth", "CFG_CustomerIDPrefix", "PC_Month", "PC_Area", "PC_Location",
               "PC_Collector", "PC_Status", "PC_DueDate", "CH_SearchText", "ST_CustomerID", "L_LocCodes", "L_Months"]:
        if nm not in wb.defined_names:
            errs.append(f"missing named range {nm}")
    for s in ("DASHBOARD", "PRINT_SHEET", "PRINT_SHEET_UR", "CUSTOMER_HISTORY", "STATEMENT", "HOME"):
        if not wb[s].protection.sheet:
            errs.append(f"{s} is not protected")
    ps = wb["PRINT_SHEET"]
    if ps.page_setup.orientation != "landscape" or ps.print_title_rows is None:
        errs.append("PRINT_SHEET print setup wrong")
    if "&P" not in (ps.oddFooter.right.text or ""):
        errs.append("PRINT_SHEET has no page numbers")
    if not wb["PRINT_SHEET_UR"].sheet_view.rightToLeft:
        errs.append("PRINT_SHEET_UR is not right-to-left")
    with zipfile.ZipFile(path) as z:
        xml = z.read("xl/workbook.xml").decode()
    if xml.count("OFFSET('PRINT_SHEET'") < 1 or "OFFSET('STATEMENT'" not in xml:
        errs.append("dynamic print areas missing")
    return errs


def recalc(path: Path) -> Path | None:
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if not soffice:
        return None
    out = Path(tempfile.mkdtemp())
    prof = Path(tempfile.mkdtemp())
    src = out / "in.xlsx"
    shutil.copy(path, src)
    subprocess.run([soffice, f"-env:UserInstallation=file://{prof}", "--headless", "--calc", "--convert-to", "xlsx",
                    "--outdir", str(out / "calc"), str(src)], capture_output=True, timeout=600)
    res = out / "calc" / "in.xlsx"
    return res if res.exists() else None


def compare(path: Path, calc_path: Path) -> list[str]:
    errs = []
    d = svc.load(Store(path, check_lock=False))
    wb = load_workbook(calc_path, data_only=True)
    # 1. no formula errors anywhere
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and c.value[:1] == "#" and c.value.rstrip("!?/0").upper() in (
                        "#N/A", "#VALUE", "#REF", "#NAME", "#DIV", "#NUM", "#NULL"):
                    errs.append(f"formula error {ws.title}!{c.coordinate} = {c.value}")
    # 2. billing rows: Excel vs Python
    ws = wb["MONTHLY_BILLING"]
    hdr = {c.value: c.column for c in ws[4]}
    xl = {}
    for r in range(5, ws.max_row + 1):
        bid = ws.cell(r, hdr["Billing ID"]).value
        if bid:
            xl[bid] = {k: ws.cell(r, hdr[k]).value for k in ("Amount Paid", "Balance", "Payment Status", "Collector ID", "Overdue")}
    for b in d.billing:
        x = xl.get(b["Billing ID"])
        if not x:
            errs.append(f"billing {b['Billing ID']} missing after recalculation")
            continue
        for k in ("Amount Paid", "Balance", "Payment Status", "Overdue"):
            if (x[k] or "") != (b[k] or "") and not (isinstance(b[k], (int, float)) and abs((x[k] or 0) - b[k]) < 0.01):
                errs.append(f"{b['Billing ID']} {k}: Excel={x[k]!r} Python={b[k]!r}")
        if (x["Collector ID"] or "") != (b["Collector ID"] or ""):
            errs.append(f"{b['Billing ID']} Collector: Excel={x['Collector ID']!r} Python={b['Collector ID']!r}")
    # 3. customers derived fields
    ws = wb["CUSTOMERS"]
    hdr = {c.value: c.column for c in ws[4]}
    for c in d.customers:
        r = c["_row"]
        for k in ("Area", "Full Address", "Monthly Fee"):
            v = ws.cell(r, hdr[k]).value
            if (v or "") != (c[k] or ""):
                errs.append(f"{c['Customer ID']} {k}: Excel={v!r} Python={c[k]!r}")
        chk = ws.cell(r, hdr["Data Check"]).value
        if chk:
            errs.append(f"{c['Customer ID']} Data Check flagged: {chk}")
    # 4. dashboard KPIs
    ws = wb["DASHBOARD"]
    month = ws["C2"].value
    s = svc.dashboard(d, month)
    expect = {"A7": s["total_customers"], "A9": s["active"], "A13": s["disconnected"], "E9": s["billing"],
              "E11": s["recovered"], "E13": s["pending"], "I7": s["paid"], "I9": s["partial"], "I13": s["pending_count"],
              "M11": s["outstanding_all"]}
    for cell, v in expect.items():
        if abs((ws[cell].value or 0) - v) > 0.01:
            errs.append(f"DASHBOARD {cell}: Excel={ws[cell].value} Python={v}")
    if abs((ws["E15"].value or 0) - s["recovery_pct"]) > 1e-6:
        errs.append("DASHBOARD recovery % differs")
    # 5. print center count & print sheet rows
    pc = wb["PRINT_CENTER"]
    f = {k: pc[c].value for k, c in (("mode", "C4"), ("month", "C5"), ("area", "C6"), ("location", "C7"),
                                     ("collector", "C8"), ("status", "C9"), ("due", "C10"))}
    bills = svc.select_bills(d, f["month"], f["area"], f["location"], f["collector"], f["status"], f["due"], mode=f["mode"])
    if pc["G4"].value != 0 and not any(pc.cell(r, 6).value for r in range(4, 104)):
        errs.append(f"PRINT_CENTER custom ID count should be 0, got {pc['G4'].value}")
    if pc["C13"].value != len(bills):
        errs.append(f"PRINT_CENTER count Excel={pc['C13'].value} Python={len(bills)}")
    ps = wb["PRINT_SHEET"]
    xl_ids = [ps.cell(r, 2).value for r in range(10, 10 + len(bills))]
    if xl_ids != [b["Customer ID"] for b in bills]:
        errs.append("PRINT_SHEET rows differ from Python selection")
    tot_row = 10 + len(bills)
    if ps.cell(tot_row, 3).value != "TOTAL" or abs((ps.cell(tot_row, 7).value or 0) - sum(b["Monthly Fee"] for b in bills)) > 0.01:
        errs.append("PRINT_SHEET total row wrong")
    # 6. cash reconciliation
    ws = wb["CASH_RECONCILIATION"]
    hdr = {c.value: c.column for c in ws[4]}
    for m in {c["Month"] for c in d.cash}:
        py = {r["Collector ID"]: r for r in svc.cash_reconciliation(d, m)}
        for r in range(5, ws.max_row + 1):
            if ws.cell(r, hdr["Month"]).value == m:
                cid = ws.cell(r, hdr["Collector ID"]).value
                if abs((ws.cell(r, hdr["Expected Collection"]).value or 0) - py[cid]["Expected Cash"]) > 0.01:
                    errs.append(f"cash {m} {cid} expected differs")
                xd = ws.cell(r, hdr["Cash Difference"]).value
                if py[cid]["Difference"] is not None and abs((xd or 0) - py[cid]["Difference"]) > 0.01:
                    errs.append(f"cash {m} {cid} difference differs")
    # 7. customer history totals
    ch = wb["CUSTOMER_HISTORY"]
    cid = ch["C7"].value
    if cid:
        h = svc.customer_history(d, cid)
        if abs((ch["H22"].value or 0) - h["total_billing"]) > 0.01 or abs((ch["H24"].value or 0) - h["total_outstanding"]) > 0.01:
            errs.append("CUSTOMER_HISTORY totals differ")
        months = [ch.cell(r, 2).value for r in range(39, 39 + len(h["bills"]))]
        if months != [b["Month"] for b in h["bills"]]:
            errs.append(f"CUSTOMER_HISTORY months differ: {months}")
    return errs


def main(path):
    path = Path(path)
    print(f"Verifying {path}")
    errs = structure(path)
    print(f"  structure: {'OK' if not errs else 'PROBLEMS'}")
    calc = recalc(path)
    if calc is None:
        print("  recalculation: skipped (LibreOffice not installed)")
    else:
        e2 = compare(path, calc)
        print(f"  recalculation vs Python: {'OK' if not e2 else 'PROBLEMS'}")
        errs += e2
    for e in errs[:50]:
        print("   - " + e)
    return 1 if errs else 0


if __name__ == "__main__":
    sys.exit(main(sys.argv[1] if len(sys.argv) > 1 else Path(__file__).resolve().parents[1] / "data" / "DEMO_CITY_LINKS_Recovery.xlsx"))
